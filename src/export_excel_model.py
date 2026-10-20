"""
Generates the comprehensive Excel Financial Model & Executive Performance Workbook.
Includes multi-tab reporting, calculated formulas, ABC Pareto, RFM summary, and What-If models.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from sqlalchemy import text
from src.db import db_session
from src.utils import setup_logger

logger = setup_logger(__name__)

def build_excel_financial_model(output_path: str = "excel/RetailIQ_Financial_Model.xlsx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    kpi_val_font = Font(name="Calibri", size=20, bold=True, color="1F4E79")
    kpi_lbl_font = Font(name="Calibri", size=10, bold=True, color="595959")

    primary_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    kpi_bg_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # -------------------------------------------------------------------------
    # TAB 1: Executive KPI Dashboard
    # -------------------------------------------------------------------------
    ws_exec = wb.create_sheet(title="Executive_Summary")
    ws_exec.views.sheetView[0].showGridLines = True

    ws_exec["A2"] = "RetailIQ Executive Performance Scorecard"
    ws_exec["A2"].font = title_font
    ws_exec["A3"] = "Enterprise Sales, Profitability & Customer Metrics (2023 - 2026)"
    ws_exec["A3"].font = subtitle_font

    # Fetch Top-level KPIs from DB
    with db_session() as conn:
        kpi_row = conn.execute(text("""
            SELECT 
                SUM(net_sales_amount) AS revenue,
                SUM(gross_profit_amount) AS profit,
                ROUND((SUM(gross_profit_amount) / NULLIF(SUM(net_sales_amount), 0)) * 100, 2) AS margin_pct,
                COUNT(DISTINCT order_id) AS total_orders,
                COUNT(DISTINCT customer_id) AS total_customers
            FROM fact_sales WHERE order_status = 'Completed';
        """)).fetchone()

    kpi_cards = [
        ("Total Net Revenue", f"${kpi_row[0]:,.2f}", "B5", "B6", "C6"),
        ("Gross Profit", f"${kpi_row[1]:,.2f}", "D5", "D6", "E6"),
        ("Gross Margin", f"{kpi_row[2]}%", "F5", "F6", "G6"),
        ("Total Orders", f"{kpi_row[3]:,}", "H5", "H6", "I6"),
        ("Active Customers", f"{kpi_row[4]:,}", "J5", "J6", "K6"),
    ]

    for title, val, top_cell, val_cell, merge_end in kpi_cards:
        ws_exec[top_cell] = title
        ws_exec[top_cell].font = kpi_lbl_font
        ws_exec[top_cell].fill = accent_fill
        ws_exec[top_cell].alignment = Alignment(horizontal="center", vertical="center")

        ws_exec[val_cell] = val
        ws_exec[val_cell].font = kpi_val_font
        ws_exec[val_cell].fill = kpi_bg_fill
        ws_exec[val_cell].alignment = Alignment(horizontal="center", vertical="center")

    # -------------------------------------------------------------------------
    # Helper for SQL export to sheets
    # -------------------------------------------------------------------------
    def export_query_to_sheet(query_file: str, sheet_name: str, max_rows: int = 500):
        with open(query_file, "r", encoding="utf-8") as f:
            sql_text = f.read()
        
        with db_session() as conn:
            df = pd.read_sql(text(sql_text), conn)

        if max_rows and len(df) > max_rows:
            df = df.head(max_rows)

        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        # Header
        headers = list(df.columns)
        ws.append([h.replace("_", " ").title() for h in headers])
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = primary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Rows
        for _, row in df.iterrows():
            ws.append(list(row))

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    logger.info("Exporting Monthly Revenue Trends...")
    export_query_to_sheet("sql/03_revenue_growth_trends.sql", "Monthly_Revenue_Trends")

    logger.info("Exporting Product ABC Pareto...")
    export_query_to_sheet("sql/04_product_abc_pareto.sql", "Product_ABC_Pareto")

    logger.info("Exporting Customer RFM Segmentation...")
    export_query_to_sheet("sql/05_rfm_customer_segmentation.sql", "Customer_RFM_Segments", max_rows=300)

    logger.info("Exporting Cohort Retention...")
    export_query_to_sheet("sql/06_cohort_retention_matrix.sql", "Cohort_Retention")

    # -------------------------------------------------------------------------
    # TAB: What-If Sensitivity Forecasting Model
    # -------------------------------------------------------------------------
    ws_whatif = wb.create_sheet(title="WhatIf_Sensitivity_Model")
    ws_whatif.views.sheetView[0].showGridLines = True

    ws_whatif["A2"] = "RetailIQ What-If Sensitivity & Price Elasticity Model"
    ws_whatif["A2"].font = title_font
    ws_whatif["A3"] = "Interactive Scenario Simulator: Price Adjustments vs Gross Profit Impact"
    ws_whatif["A3"].font = subtitle_font

    # Parameters
    params = [
        ("Base Annual Sales Volume (Units)", 35000, "B6", "#,##0"),
        ("Baseline Selling Price ($)", 250.00, "B7", "$#,##0.00"),
        ("Baseline Unit Cost / COGS ($)", 175.00, "B8", "$#,##0.00"),
        ("Price Elasticity of Demand (Ped)", -1.25, "B9", "0.00"),
        ("Scenario Price Change %", 0.05, "B10", "0.0%"),
    ]

    for lbl, val, cell_ref, num_fmt in params:
        row_num = int(cell_ref[1:])
        ws_whatif.cell(row=row_num, column=1, value=lbl).font = Font(name="Calibri", bold=True)
        c = ws_whatif[cell_ref]
        c.value = val
        c.number_format = num_fmt
        c.fill = accent_fill

    # Projected Outputs (Formulas)
    ws_whatif["A13"] = "Projected New Selling Price ($)"
    ws_whatif["B13"] = "=B7*(1+B10)"
    ws_whatif["B13"].number_format = "$#,##0.00"

    ws_whatif["A14"] = "Projected Sales Volume Change %"
    ws_whatif["B14"] = "=B10*B9"
    ws_whatif["B14"].number_format = "0.0%"

    ws_whatif["A15"] = "Projected Annual Units Sold"
    ws_whatif["B15"] = "=B6*(1+B14)"
    ws_whatif["B15"].number_format = "#,##0"

    ws_whatif["A16"] = "Projected Total Net Revenue ($)"
    ws_whatif["B16"] = "=B15*B13"
    ws_whatif["B16"].number_format = "$#,##0.00"

    ws_whatif["A17"] = "Projected Total COGS ($)"
    ws_whatif["B17"] = "=B15*B8"
    ws_whatif["B17"].number_format = "$#,##0.00"

    ws_whatif["A18"] = "Projected Gross Profit ($)"
    ws_whatif["B18"] = "=B16-B17"
    ws_whatif["B18"].number_format = "$#,##0.00"

    ws_whatif["A19"] = "Projected Gross Margin %"
    ws_whatif["B19"] = "=B18/B16"
    ws_whatif["B19"].number_format = "0.00%"

    for r in range(13, 20):
        ws_whatif.cell(row=r, column=1).font = Font(name="Calibri", bold=True)
        ws_whatif.cell(row=r, column=2).font = Font(name="Calibri", bold=True, color="1F4E79")
        ws_whatif.cell(row=r, column=2).fill = kpi_bg_fill

    # Set column widths
    ws_whatif.column_dimensions["A"].width = 38
    ws_whatif.column_dimensions["B"].width = 22

    wb.save(output_path)
    logger.info(f"Successfully generated Excel Financial Model -> {output_path}")

if __name__ == "__main__":
    build_excel_financial_model()
